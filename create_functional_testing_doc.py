from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\xampp\htdocs\BloodLink\BloodLink Functional Testing Form.docx"

modules = [
    ("1. Authentication and Role Based Access", [
        ("FR1.1", "Log in with valid seeded credentials for each authorized role", "System authenticates the user and opens the correct role dashboard."),
        ("FR1.2", "Log in using an incorrect password", "System denies access and displays a login error message."),
        ("FR1.3", "Log in using an unregistered email address", "System denies access and displays an account-not-found or login error message."),
        ("FR1.4", "Submit login with required fields blank", "System identifies the required email and password fields."),
        ("FR1.5", "Open a protected dashboard without an authenticated session", "System prevents unauthorized access to the protected dashboard."),
        ("FR1.6", "Log out from a portal", "System ends the session and returns the user to the landing or login page."),
        ("FR1.7", "Verify role restricted modules", "User sees only modules and actions assigned to their role."),
    ]),
    ("2. Donor Management Unit", [
        ("FR2.1", "Register a donor with complete and valid required information", "New donor record is saved and appears in donor management."),
        ("FR2.2", "Register a donor with incomplete required information", "System prevents submission and identifies the missing fields."),
        ("FR2.3", "Search an existing donor by name or donor identifier", "Relevant donor record is displayed."),
        ("FR2.4", "View an existing donor profile", "Correct demographic, blood type, and donation information is displayed."),
        ("FR2.5", "Update a donor profile", "Updated donor information is saved and displayed correctly."),
        ("FR2.6", "Record a donation with valid donor and collection details", "Donation record is saved and linked to the correct donor."),
        ("FR2.7", "Record a donation with missing required information", "System prevents submission and identifies missing fields."),
        ("FR2.8", "View donation records", "Donation records display the correct donor, serial number, date, and status."),
        ("FR2.9", "Search or filter donation records", "Only records matching the selected search or filter criteria are displayed."),
    ]),
    ("3. Serology and Transmissible Disease Screening", [
        ("FR3.1", "View donations awaiting laboratory results", "Pending donations are displayed for serology processing."),
        ("FR3.2", "Open a pending donation from its serial number", "Laboratory result form opens with the selected donation details."),
        ("FR3.3", "Encode complete non-reactive screening results", "Lab result is saved and the donation is marked Accepted or Cleared according to the outcome."),
        ("FR3.4", "Encode reactive or indeterminate test results", "System records the result and applies the appropriate screening outcome or quarantine status."),
        ("FR3.5", "Submit a laboratory result with missing required fields", "System prevents submission and identifies required fields."),
        ("FR3.6", "Search or filter laboratory results", "Only results matching the selected serial number or result filter are displayed."),
        ("FR3.7", "View an encoded laboratory result", "Saved blood type, TTI results, outcome, and remarks are displayed correctly."),
    ]),
    ("4. Blood Component Production and Inventory", [
        ("FR4.1", "View serology-cleared donations awaiting component processing", "Eligible donations are listed for processing."),
        ("FR4.2", "Open component processing using a cleared donation serial number", "Processing form opens with the selected serial number."),
        ("FR4.3", "Record one or more blood components with valid details", "Component units are saved and linked to the correct donation."),
        ("FR4.4", "Attempt to process a donation without a laboratory result", "System prevents processing and informs the user that laboratory results are required."),
        ("FR4.5", "Attempt to process a deferred or non-accepted donation", "System prevents component processing for the ineligible donation."),
        ("FR4.6", "Enter a component volume outside the acceptable range", "System flags the entry and applies the required safety handling or validation."),
        ("FR4.7", "View blood component inventory ledger", "Recorded units display the correct serial number, blood type, component, and status."),
        ("FR4.8", "Search or filter inventory by component or status", "Only matching inventory units are displayed."),
        ("FR4.9", "View units approaching expiration", "Expiring units are correctly identified in the inventory indicators."),
    ]),
    ("5. Hospital Blood Requests and Issuance", [
        ("FR5.1", "Hospital user creates a blood request with complete valid information", "Request is submitted, assigned a reference number, and appears in the hospital request list."),
        ("FR5.2", "Hospital user submits a request with required fields missing", "System prevents submission and identifies the missing information."),
        ("FR5.3", "Hospital user views submitted blood requests", "Only the user hospital's requests and their current statuses are displayed."),
        ("FR5.4", "Issuance Personnel views incoming blood requests", "Request queue displays submitted hospital requests with correct details and status."),
        ("FR5.5", "Verify a submitted blood request", "Request status updates to Verified when stock and request details are validated."),
        ("FR5.6", "Reject an ineligible or incomplete request", "System records the rejection and displays the updated request status."),
        ("FR5.7", "Process an approved blood request", "System records the issuance workflow and updates request processing status."),
        ("FR5.8", "Approve blood release", "Release status is recorded and issued units are reflected in the issuance record."),
        ("FR5.9", "Record a walk-in issuance", "Walk-in issuance is saved with correct unit, recipient, and release details."),
        ("FR5.10", "View issuance audit log", "Audit log displays completed issuance and release details correctly."),
    ]),
    ("6. Donor Recall and Notifications", [
        ("FR6.1", "View donors eligible for recall based on the donation interval", "System identifies and lists eligible repeat donors."),
        ("FR6.2", "Exclude a donor not yet eligible for recall", "Ineligible donor is not included in the eligible recall list."),
        ("FR6.3", "Generate a donor recall message", "System creates a reminder message for the selected eligible donor."),
        ("FR6.4", "Send a recall reminder through PhilSMS", "System sends the notification request and records the delivery or sending status."),
        ("FR6.5", "View donor recall history", "Recall records show donor, message date, and notification status."),
    ]),
    ("7. Demand Forecasting and Equity Based Distribution", [
        ("FR7.1", "Generate a Multiple Linear Regression demand forecast using available historical distribution data", "System produces forecast values by blood type for the selected forecast period."),
        ("FR7.2", "Filter demand forecast by blood type, hospital, or forecast period", "Forecast display updates to the selected criteria."),
        ("FR7.3", "View demand versus inventory gap analysis", "System displays predicted demand, available inventory, gap, and shortage or surplus indication."),
        ("FR7.4", "Generate equity based distribution recommendations", "System recommends proportional component allocation using predicted demand and available inventory."),
        ("FR7.5", "View distribution recommendation details for a hospital", "Recommended units, basis of allocation, and inventory context are displayed."),
        ("FR7.6", "Generate recommendations when available inventory is insufficient", "System returns proportional or constrained recommendations without allocating more units than available."),
    ]),
    ("8. Administration and Audit", [
        ("FR8.1", "Super Admin views the user management list", "System displays user accounts with correct role and account details."),
        ("FR8.2", "Create a user account with a valid assigned role", "New user account is saved and assigned the selected role."),
        ("FR8.3", "Create a user account with duplicate email", "System prevents duplicate account creation and displays a validation message."),
        ("FR8.4", "Update an existing user account", "Updated user information and role are saved correctly."),
        ("FR8.5", "Deactivate or reactivate a user account", "Account status changes and access follows the selected status."),
        ("FR8.6", "View audit activity or system records", "Relevant activity records display the user, action, and timestamp correctly."),
    ]),
]

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9')

def set_cell(cell, text, bold=False, size=8.2, color='000000', align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text); r.bold = bold; r.font.name = 'Arial'; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    r._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(cell)

def set_width(cell, inches):
    cell.width = Inches(inches)
    tcPr = cell._tc.get_or_add_tcPr(); tcW = tcPr.find(qn('w:tcW'))
    if tcW is not None: tcW.set(qn('w:w'), str(int(inches * 1440))); tcW.set(qn('w:type'), 'dxa')

def add_table(doc, cases):
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.62, 3.15, 3.15, 2.0, 1.55]
    headers = ['FR#', 'Description', 'Expected Result', 'Actual Result', 'Remarks']
    for i, label in enumerate(headers):
        set_width(table.rows[0].cells[i], widths[i]); shade(table.rows[0].cells[i], '1F4E78')
        set_cell(table.rows[0].cells[i], label, True, 8.2, 'FFFFFF', WD_ALIGN_PARAGRAPH.CENTER)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for n, (code, desc, expected) in enumerate(cases):
        cells = table.add_row().cells
        fill = 'F6F9FC' if n % 2 else 'FFFFFF'
        values = [code, desc, expected, '', '']
        for i, value in enumerate(values):
            set_width(cells[i], widths[i]); shade(cells[i], fill)
            set_cell(cells[i], value, i == 0, 8.0, '000000', WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Inches(0.55); section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.48); section.right_margin = Inches(0.48)

styles = doc.styles
styles['Normal'].font.name = 'Arial'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); styles['Normal'].font.size = Pt(9)
for name in ['Title', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'Arial'; styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); styles[name].font.color.rgb = RGBColor(0,0,0)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('BloodLink Functional Testing Form'); run.font.name = 'Arial'; run.font.size = Pt(20); run.font.color.rgb = RGBColor(0,0,0)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Blood Inventory Monitoring and Decision Support System'); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(10)
doc.add_paragraph('This form documents functional test cases for BloodLink. Execute each case using the stated scenario, record the observed outcome in Actual Result, and note evidence, defects, or follow-up actions in Remarks.').paragraph_format.space_after = Pt(8)

info = doc.add_table(rows=2, cols=4); info.alignment = WD_TABLE_ALIGNMENT.CENTER; info.autofit = False
labels = [('Project', 'BloodLink'), ('Testing Type', 'Functional Testing'), ('Tester', '________________________'), ('Date', '________________________'), ('Build or Version', '________________________'), ('Test Environment', '________________________'), ('Result Codes', 'P Passed   F Failed   PEND Pending'), ('Evidence Reference', '________________________')]
for idx, (label, value) in enumerate(labels):
    cell = info.rows[idx//4].cells[idx%4]; shade(cell, 'F6F9FC'); borders(cell)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    a = p.add_run(label + ': '); a.bold = True; a.font.name = 'Arial'; a.font.size = Pt(8.5)
    b = p.add_run(value); b.font.name = 'Arial'; b.font.size = Pt(8.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
doc.add_paragraph().paragraph_format.space_after = Pt(4)

for index, (heading, cases) in enumerate(modules):
    h = doc.add_paragraph(style='Heading 1'); h.paragraph_format.space_before = Pt(8); h.paragraph_format.space_after = Pt(4)
    rr = h.add_run(heading); rr.font.name = 'Arial'; rr.font.size = Pt(12); rr.font.color.rgb = RGBColor(0,0,0)
    add_table(doc, cases)

h = doc.add_paragraph(style='Heading 1'); h.add_run('Functional Testing Summary').font.color.rgb = RGBColor(0,0,0)
doc.add_paragraph('Complete this summary after executing the individual functional test cases. Keep Pending separate from Failed until a final result is established.')
summary = doc.add_table(rows=1, cols=6); summary.alignment = WD_TABLE_ALIGNMENT.CENTER; summary.autofit = False
for i, label in enumerate(['Module', 'Passed', 'Failed', 'Pending', 'Total', 'Pass Rate']):
    set_width(summary.rows[0].cells[i], [5.2, 1.0, 1.0, 1.0, 1.0, 1.2][i]); shade(summary.rows[0].cells[i], '1F4E78'); set_cell(summary.rows[0].cells[i], label, True, 8.5, 'FFFFFF', WD_ALIGN_PARAGRAPH.CENTER)
for heading, cases in modules:
    cells = summary.add_row().cells
    values = [heading.split('. ', 1)[1], '', '', '', str(len(cases)), '']
    for i, val in enumerate(values):
        set_width(cells[i], [5.2, 1.0, 1.0, 1.0, 1.0, 1.2][i]); shade(cells[i], 'FFFFFF'); set_cell(cells[i], val, False, 8.2, '000000', WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
total = sum(len(cases) for _, cases in modules)
cells = summary.add_row().cells
for i, val in enumerate(['TOTAL', '', '', '', str(total), '']):
    set_width(cells[i], [5.2, 1.0, 1.0, 1.0, 1.0, 1.2][i]); shade(cells[i], 'DCE6F1'); set_cell(cells[i], val, True, 8.4, '000000', WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
p = doc.add_paragraph(); p.add_run('Suggested computation: ').bold = True; p.add_run('Total Test Cases = Passed + Failed + Pending. Pass Rate = Passed divided by (Passed + Failed) multiplied by 100. Do not include Pending in the Pass Rate until it has been executed.')

footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('BloodLink Functional Testing Form'); fr.font.name = 'Arial'; fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(100,100,100)

doc.core_properties.title = 'BloodLink Functional Testing Form'
doc.core_properties.author = 'BloodLink Project Team'
doc.save(OUT)
print(OUT)
